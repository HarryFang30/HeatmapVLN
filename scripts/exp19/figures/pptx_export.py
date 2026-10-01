"""PowerPoint copy of a finished paper figure, for editing in PowerPoint / WPS / Keynote.

One slide the size of the figure.  Every text is a text box; the box a text is drawn on (the circles of the
key-moment numbers, the white box behind a route-map label) is a shape right under it; everything else is one
picture per axes (PNG, ``DPI``, transparent outside the axes' own marks; inset axes go with their parent) at the
same place.  A panel's picture and its texts (tick labels included) form one group, so a panel moves as one; texts
given the same key with ``same_box`` (a heading's style runs) become one text box.  No notes page: python-pptx's
notes slide hangs the macOS QuickLook preview (the caption has its own ``.txt``).

Placement: each text line's baseline and start are where matplotlib drew them (recorded from the Agg renderer).  Text
boxes have no insets and "exactly" line spacing equal to the line's ascent + descent (Arial 1.117 em, Microsoft YaHei
1.320 em); with that spacing every layout rule the Office-format editors use puts the baseline one ascent below the
top of the line, so a box's top is its first baseline minus the ascent.  Lines further apart than that get the rest
as space before the paragraph; lines closer than that (stacked CJK characters, a CJK title over a Latin line) give up
the difference from both line heights: exact where an editor centres a line's box on its glyphs, within ~0.25 pt
where it puts the baseline at 0.8 of the line (what PowerPoint's PDF export suggests).  A box is ``SLACK`` wider
than its text on the side away from its alignment, within the page, so an editor that wraps anyway (the macOS
QuickLook preview ignores "do not wrap") keeps every line whole.  Fonts: Nimbus Sans -> Arial (both have
Helvetica's widths), CJK -> Microsoft YaHei (installed with Office on Windows and macOS).

Texts drawn as paths (path effects) stay in the pictures, and figure-level artists other than texts become one
picture above the panels; ``remarks`` in the result says when either happens.
"""
from __future__ import annotations

import io
import re
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np

import matplotlib.colors as mcolors
from matplotlib.backends.backend_agg import FigureCanvasAgg, RendererAgg
from matplotlib.font_manager import weight_dict
from matplotlib.patches import BoxStyle
from matplotlib.text import Annotation, Text

EMU_PER_PT = 12700
DPI = 600  # pictures
LATIN_FONT = "Arial"
CJK_FONT = "Microsoft YaHei"
METRICS = {LATIN_FONT: (0.9052734375, 0.2119140625), CJK_FONT: (1.05810546875, 0.26171875)}  # (ascent, descent) em
SPACE_EM = 570 / 2048  # Arial's space (Nimbus Sans: 0.278)
BOX_KEY = "_pptx_box"  # attribute ``same_box`` sets on a Text (not its gid: an SVG id has to be unique)
SLACK = (0.06, 2.0)  # extra box width: fraction of the text width + pt
_CJK = re.compile(r"[⺀-鿿豈-﫿︰-﹏＀-￯　-〿]")
_MATH = ((r"\tilde{\mathrm{Z}}", "Z̃"), (r"\tilde{Z}", "Z̃"), (r"^\circ", "°"), (r"\circ", "°"),
         (r"\sigma", "σ"), (r"\times", "×"), (r"\pm", "±"), (r"\mathrm", ""), ("{", ""), ("}", ""), ("$", ""))


@dataclass
class Run:
    text: str
    size: float  # pt
    bold: bool
    italic: bool
    rgba: Tuple[float, float, float, float]
    x: float  # pen position at the start of the baseline, pt from the page's top-left corner
    y: float
    w: float  # advance width, pt

    @property
    def cjk(self) -> bool:
        return bool(_CJK.search(self.text))

    def ascent(self) -> float:
        return max(METRICS[LATIN_FONT][0], METRICS[CJK_FONT][0] if self.cjk else 0.0) * self.size

    def descent(self) -> float:
        return max(METRICS[LATIN_FONT][1], METRICS[CJK_FONT][1] if self.cjk else 0.0) * self.size


@dataclass
class Item:
    """One text box: lines of runs (left to right), all at ``angle`` (degrees, counter-clockwise)."""
    lines: List[List[Run]]
    angle: float
    align: str  # "l", "ctr", "r"
    parent: object  # the axes, or None for a figure text
    boxes: List[dict] = field(default_factory=list)  # drawn behind: kind, left, top, w, h (pt), fill, edge, lw, radius

    @property
    def text(self) -> str:
        return "\n".join(" ".join(r.text for r in line) for line in self.lines)


def plain(s: str) -> str:
    """A mathtext string as plain text (the few commands these figures use)."""
    for a, b in _MATH:
        s = s.replace(a, b)
    return s


# --------------------------------------------------------------------------- #
# Recording what matplotlib draws
# --------------------------------------------------------------------------- #
class Recorded(list):
    """(Text, lines) per drawn text; ``as_paths``: visible texts that drew no line (path effects draw paths)."""

    def __init__(self):
        super().__init__()
        self.as_paths: List[Text] = []


@contextmanager
def recording():
    """Record, per drawn ``Text``, every line Agg draws: (x, y from the top, string, font properties, angle, ismath,
    width), in pixels."""
    stack: List[list] = []
    drawn = Recorded()
    text_draw, draw_text = Text.draw, RendererAgg.draw_text

    def rec_text_draw(self, renderer):
        stack.append([])
        try:
            return text_draw(self, renderer)
        finally:
            lines = stack.pop()
            if lines:
                drawn.append((self, lines))
            elif self.get_visible() and self.get_text().strip() and self.get_path_effects():
                drawn.as_paths.append(self)

    def rec_draw_text(self, gc, x, y, s, prop, angle, ismath=False, mtext=None):
        if stack:
            w = self.get_text_width_height_descent(s, prop, ismath)[0]
            stack[-1].append((float(x), float(y), s, prop.copy(), float(angle), ismath, float(w)))
        return draw_text(self, gc, x, y, s, prop, angle, ismath=ismath, mtext=mtext)

    Text.draw, RendererAgg.draw_text = rec_text_draw, rec_draw_text
    try:
        yield drawn
    finally:
        Text.draw, RendererAgg.draw_text = text_draw, draw_text


@contextmanager
def agg_at(fig, dpi: float):
    """The figure on an Agg canvas at ``dpi`` with no page background; everything restored afterwards."""
    canvas, old_dpi, bg = fig.canvas, fig.dpi, fig.patch.get_visible()
    agg = FigureCanvasAgg(fig)
    fig.dpi = dpi
    fig.patch.set_visible(False)
    try:
        yield agg
    finally:
        fig.patch.set_visible(bg)
        fig.dpi = old_dpi
        fig.set_canvas(canvas)


def _is_bold(weight) -> bool:
    w = weight if isinstance(weight, (int, float)) else weight_dict.get(str(weight), 400)
    return w >= 600


ROUND_STYLES = tuple(c for c in (getattr(BoxStyle, n, None) for n in ("Circle", "Ellipse")) if c)


def _patch_info(t: Text, renderer, px2pt: float, height_px: float, notes: List[str]) -> Optional[dict]:
    p = t.get_bbox_patch()
    if p is None or not p.get_visible():
        return None
    ext = p.get_window_extent(renderer)
    style = p.get_boxstyle()
    kind, radius = "rect", 0.0
    if isinstance(style, ROUND_STYLES):
        kind = "ellipse"
    elif isinstance(style, BoxStyle.Round):
        rs = style.rounding_size if style.rounding_size is not None else style.pad
        radius = rs * p.get_mutation_scale() * px2pt
        kind = "roundRect" if radius > 0 else "rect"
    elif not isinstance(style, BoxStyle.Square):
        notes.append(f"box behind {t.get_text()!r} ({type(style).__name__}) written as a rectangle")
    return {"kind": kind, "left": ext.x0 * px2pt, "top": (height_px - ext.y1) * px2pt, "w": ext.width * px2pt,
            "h": ext.height * px2pt, "fill": tuple(p.get_facecolor()), "edge": tuple(p.get_edgecolor()),
            "lw": float(p.get_linewidth()), "radius": radius}


def axis_owners(fig) -> Dict[int, object]:
    """Tick labels, axis labels and offset texts have no ``.axes``: the axes whose x / y axis drew them, by id."""
    return {id(t): ax for ax in fig.axes for axis in (ax.xaxis, ax.yaxis) for t in axis.findobj(Text)}


def top_axes(fig) -> Dict[int, object]:
    """Every axes, inset (child) axes included, by id -> the figure-level axes that draws it."""
    out: Dict[int, object] = {}

    def walk(ax, root):
        out[id(ax)] = root
        for child in getattr(ax, "child_axes", []):
            walk(child, root)
    for ax in fig.axes:
        walk(ax, ax)
    return out


def same_box(text: Text, key: str) -> Text:
    """Texts given the same ``key`` become one text box (lines by baseline, style runs left to right)."""
    setattr(text, BOX_KEY, key)
    return text


def texts_as_items(drawn, renderer, dpi: float, height_px: float, owners: Optional[Dict[int, object]] = None,
                   tops: Optional[Dict[int, object]] = None) -> Tuple[List[Item], List[str]]:
    """Items from the recorded texts, runs with the same ``same_box`` key merged into one item.  An item's parent is
    the figure-level axes (``tops``) of the text's axes, or of the axes in ``owners`` that drew it; None (the page)
    for a figure text."""
    px2pt = 72.0 / dpi
    items, merge, notes = [], {}, []
    owners, tops = owners or {}, tops or {}
    for t, lines in drawn:
        ax = t.axes if t.axes is not None else owners.get(id(t))
        parent = tops.get(id(ax)) if ax is not None else None
        rgba = mcolors.to_rgba(t.get_color(), t.get_alpha())
        runs = []
        for x, y, s, prop, angle, ismath, w in lines:
            if ismath:
                notes.append(f"mathtext {s!r} written as {plain(s)!r}")
                s = plain(s)
            if not s.strip():
                continue
            runs.append(Run(s, prop.get_size_in_points(), _is_bold(prop.get_weight()),
                            prop.get_style() in ("italic", "oblique"), rgba, x * px2pt, y * px2pt, w * px2pt))
        if not runs or rgba[3] == 0:
            continue
        angle = lines[0][4] % 360.0
        patch = _patch_info(t, renderer, px2pt, height_px, notes)
        key = getattr(t, BOX_KEY, None)
        if key is not None and len(runs) == 1 and angle == 0:
            merge.setdefault((key, id(parent)), []).append((runs[0], parent, patch))
            continue
        if len(runs) == 1 and angle == 0:
            align = {"left": "l", "center": "ctr", "right": "r"}[t.get_horizontalalignment()]
        elif len(runs) > 1:
            align = {"left": "l", "center": "ctr", "right": "r"}[t._get_multialignment()]
        else:
            align = "ctr"  # a rotated line: keep its centre
        if patch is not None and angle % 90 != 0:
            notes.append(f"box behind {t.get_text()!r} dropped (rotated {angle:g}°)")
            patch = None
        items.append(Item([[r] for r in runs], angle, align, parent, [patch] if patch else []))
    for _, pieces in merge.items():  # runs centred on one line sit a little apart: one line within half a size
        lines_: List[List[Run]] = []
        for r in sorted((r for r, _, _ in pieces), key=lambda r: r.y):
            if lines_ and r.y - baseline(lines_[-1]) <= 0.5 * max(r.size, *(q.size for q in lines_[-1])):
                lines_[-1].append(r)
            else:
                lines_.append([r])
        items.append(Item([sorted(ln, key=lambda r: r.x) for ln in lines_], 0.0, "l", pieces[0][1],
                          [pb for _, _, pb in pieces if pb]))
    return items, notes


# --------------------------------------------------------------------------- #
# Geometry
# --------------------------------------------------------------------------- #
def baseline(line: Sequence[Run]) -> float:
    """A line's baseline: its largest run's (runs merged from texts centred on one line differ a little)."""
    return max(line, key=lambda r: r.size).y


def line_metrics(line: Sequence[Run]) -> Tuple[float, float]:
    return max(r.ascent() for r in line), max(r.descent() for r in line)


def layout(item: Item, page_w: Optional[float] = None) -> dict:
    """The text box of an item: left, top, w, h (pt, the unrotated box), rotation (clockwise) and per line its
    "exactly" spacing, space before and left margin (pt).  ``page_w``: keep an upright box's slack on the page."""
    th = np.radians(item.angle)
    eu, ev = np.array([np.cos(th), -np.sin(th)]), np.array([np.sin(th), np.cos(th)])  # along / below the text
    p0 = np.array([item.lines[0][0].x, baseline(item.lines[0])])
    starts = [np.array([ln[0].x, baseline(ln)]) - p0 for ln in item.lines]
    u = [float(s @ eu) for s in starts]
    v = [float(s @ ev) for s in starts]
    widths = [ln[-1].x - ln[0].x + ln[-1].w if item.angle == 0 else sum(r.w for r in ln) for ln in item.lines]
    mets = [line_metrics(ln) for ln in item.lines]
    n = len(mets)
    # short[i]: how much closer lines i-1 and i are than descent + ascent; the two line heights give up 2 x that
    # between them (cut), so the boxes stay centred on the glyphs; what a cut leaves over is space before.
    short = [0.0] + [mets[i - 1][1] + mets[i][0] - (v[i] - v[i - 1]) for i in range(1, n)]
    cut = [0.0] * n
    cut[0] = max(short[1], 0.0) if n > 1 else 0.0
    lines = []
    for i, (a, d) in enumerate(mets):
        before = 0.0
        if i:
            cut[i] = max(2 * short[i] - cut[i - 1], 0.0) if short[i] > 0 else 0.0
            before = max((cut[i - 1] + cut[i]) / 2 - short[i], 0.0)
        lines.append({"spacing": a + d - cut[i], "before": before})
    top = v[0] - mets[0][0] + cut[0] / 2
    bottom = v[-1] + mets[-1][1] - cut[-1] / 2
    box = _hbox(item.align, u, widths, p0[0] if item.angle == 0 else None, page_w)
    for ln, off in zip(lines, box["margins"]):
        ln["margin"] = off
    return _place(box["u0"], box["w"], top, bottom - top, eu, ev, p0, item.angle, lines)


def _hbox(align: str, u: Sequence[float], widths: Sequence[float], x0: Optional[float] = None,
          page_w: Optional[float] = None) -> dict:
    """Box left and width along the text (with ``SLACK``; for an upright box at page x ``x0``, not past the page
    edges ``page_w`` where the text itself is inside), and per line a left margin (left-aligned boxes only)."""
    lo = min(u)
    hi = max(ui + w for ui, w in zip(u, widths))
    slack = SLACK[0] * (hi - lo) + SLACK[1]
    clip = x0 is not None and page_w is not None
    if align == "l":
        right = hi + slack
        if clip:
            right = max(min(right, page_w - x0), hi)
        return {"u0": lo, "w": right - lo, "margins": [ui - lo for ui in u]}
    if align == "r":
        left = lo - slack
        if clip:
            left = min(max(left, -x0), lo)
        return {"u0": left, "w": hi - left, "margins": [0.0] * len(u)}
    c = u[0] + widths[0] / 2
    need = max(max(abs(ui - c), abs(ui + w - c)) for ui, w in zip(u, widths))
    half = need + slack / 2
    if clip:
        half = max(min(half, x0 + c, page_w - x0 - c), need)
    return {"u0": c - half, "w": 2 * half, "margins": [0.0] * len(u)}


def _place(u0, w, v0, h, eu, ev, p0, angle, lines) -> dict:
    centre = p0 + (u0 + w / 2) * eu + (v0 + h / 2) * ev
    return {"left": float(centre[0] - w / 2), "top": float(centre[1] - h / 2), "w": max(w, 0.5), "h": max(h, 0.5),
            "rotation": (-angle) % 360.0, "lines": lines}


# --------------------------------------------------------------------------- #
# Pictures
# --------------------------------------------------------------------------- #
def render_rgba(canvas) -> np.ndarray:
    canvas.draw()
    return np.asarray(canvas.buffer_rgba()).copy()


def crop(rgba: np.ndarray) -> Optional[Tuple[int, int, int, int]]:
    """(x0, y0, x1, y1) of the pixels with any opacity, None when there are none."""
    alpha = rgba[..., 3] > 0
    rows, cols = np.flatnonzero(alpha.any(axis=1)), np.flatnonzero(alpha.any(axis=0))
    if not len(rows):
        return None
    return int(cols[0]), int(rows[0]), int(cols[-1]) + 1, int(rows[-1]) + 1


def png_bytes(rgba: np.ndarray, dpi: float) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.fromarray(np.ascontiguousarray(rgba)).save(buf, format="PNG", dpi=(dpi, dpi))
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PowerPoint
# --------------------------------------------------------------------------- #
def _emu(pt: float) -> int:
    return int(round(pt * EMU_PER_PT))


def _hex(rgba) -> str:
    return mcolors.to_hex(rgba[:3], keep_alpha=False)[1:].upper()


def _solid(parent, rgba) -> None:
    """``<a:srgbClr>`` (with ``<a:alpha>`` when translucent) under ``parent`` (an ``<a:solidFill>``)."""
    from pptx.oxml.ns import qn

    clr = parent.makeelement(qn("a:srgbClr"), {"val": _hex(rgba)})
    if rgba[3] < 0.999:
        clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(round(rgba[3] * 100000)))}))
    parent.append(clr)


def _run_props(rPr, r: Run, *, spc: Optional[float] = None) -> None:
    """Size, weight, style, colour, fonts and language of a run (``rPr`` an ``<a:rPr>`` / ``<a:endParaRPr>``)."""
    from pptx.oxml.ns import qn

    for child in list(rPr):
        rPr.remove(child)
    rPr.set("lang", "zh-CN" if r.cjk else "en-US")
    rPr.set("altLang", "en-US")
    rPr.set("sz", str(int(round(r.size * 100))))
    rPr.set("b", "1" if r.bold else "0")
    rPr.set("i", "1" if r.italic else "0")
    rPr.set("noProof", "1")
    rPr.set("dirty", "0")
    if spc is not None:
        rPr.set("spc", str(int(round(spc * 100))))
    fill = rPr.makeelement(qn("a:solidFill"), {})
    _solid(fill, r.rgba)
    rPr.append(fill)
    for tag, face in (("a:latin", LATIN_FONT), ("a:ea", CJK_FONT), ("a:cs", LATIN_FONT)):
        rPr.append(rPr.makeelement(qn(tag), {"typeface": face}))


def _fill_text(tf, item: Item, geo: dict) -> None:
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.oxml.ns import qn

    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, (line, ln) in enumerate(zip(item.lines, geo["lines"])):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("algn", item.align)
        pPr.set("marL", str(_emu(ln["margin"])))
        pPr.set("indent", "0")
        for tag, val in (("a:lnSpc", ln["spacing"]), ("a:spcBef", ln["before"]), ("a:spcAft", 0.0)):
            el = pPr.makeelement(qn(tag), {})
            el.append(el.makeelement(qn("a:spcPts"), {"val": str(int(round(val * 100)))}))
            pPr.append(el)
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        for j, r in enumerate(line):
            if j:  # the gap between two style runs: a space widened (or narrowed) to it
                gap = r.x - (line[j - 1].x + line[j - 1].w)
                sep = p.add_run()
                sep.text = " "
                _run_props(sep._r.get_or_add_rPr(), r, spc=gap - SPACE_EM * r.size)
            run = p.add_run()
            run.text = r.text
            _run_props(run._r.get_or_add_rPr(), r)
        _run_props(p._p.get_or_add_endParaRPr(), line[-1])


def _strip_style(shape) -> None:
    """Drop the theme style reference python-pptx gives an auto shape (theme line, effects, font colour)."""
    from pptx.oxml.ns import qn

    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def add_box(shapes, pb: dict, name: str):
    """The box a text is drawn on, as a shape without text."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt

    kind = {"ellipse": MSO_SHAPE.OVAL, "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE}.get(pb["kind"], MSO_SHAPE.RECTANGLE)
    sh = shapes.add_shape(kind, _emu(pb["left"]), _emu(pb["top"]), _emu(pb["w"]), _emu(pb["h"]))
    _strip_style(sh)
    if pb["kind"] == "roundRect":
        sh.adjustments[0] = min(pb["radius"] / max(min(pb["w"], pb["h"]), 1e-6), 0.5)
    if pb["fill"][3] > 0:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(_hex(pb["fill"]))
        if pb["fill"][3] < 0.999:
            _set_alpha(sh.fill._xPr, pb["fill"][3])
    else:
        sh.fill.background()
    if pb["edge"][3] > 0 and pb["lw"] > 0:
        sh.line.width = Pt(pb["lw"])
        sh.line.color.rgb = RGBColor.from_string(_hex(pb["edge"]))
    else:
        sh.line.fill.background()
    sh.name = name
    return sh


def add_item(shapes, item: Item, name: str, page_w: Optional[float] = None):
    """The item's boxes (if any), then its text box."""
    for pb in item.boxes:
        add_box(shapes, pb, f"box {name}")
    geo = layout(item, page_w)
    sh = shapes.add_textbox(_emu(geo["left"]), _emu(geo["top"]), _emu(geo["w"]), _emu(geo["h"]))
    if geo["rotation"]:
        sh.rotation = geo["rotation"]
    _fill_text(sh.text_frame, item, geo)
    sh.name = name
    return sh


def _set_alpha(spPr, alpha: float) -> None:
    from pptx.oxml.ns import qn

    clr = spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(round(alpha * 100000)))}))


def _short(text: str, n: int = 40) -> str:
    s = " ".join(text.split())
    return s if len(s) <= n else s[: n - 1] + "…"


@contextmanager
def texts_hidden(texts: Sequence[Text]):
    """The texts out of the pictures.  An annotation with an arrow keeps the arrow: its text turns transparent and its
    box hidden instead (a hidden annotation draws nothing)."""
    saved = []
    for t in texts:
        if isinstance(t, Annotation) and t.arrow_patch is not None:
            box = t.get_bbox_patch()
            saved.append((t, "alpha", t.get_alpha(), box.get_visible() if box is not None else None))
            t.set_alpha(0.0)
            if box is not None:
                box.set_visible(False)
        else:
            saved.append((t, "visible", t.get_visible(), None))
            t.set_visible(False)
    try:
        yield
    finally:
        for t, how, value, box_vis in saved:
            if how == "alpha":
                t.set_alpha(value)
                if box_vis is not None:
                    t.get_bbox_patch().set_visible(box_vis)
            else:
                t.set_visible(value)


@contextmanager
def frozen_layout(fig):
    """No layout engine between the draws (positions as the last draw left them), the figure's engine back after."""
    engine = fig.get_layout_engine()
    if engine is not None:
        fig.set_layout_engine("none")
    try:
        yield
    finally:
        if engine is not None:
            fig.set_layout_engine(engine)


def _thumbnail(rgba: np.ndarray, width: int = 256) -> bytes:
    """The figure on white as a small JPEG (the file's preview icon)."""
    from PIL import Image

    im = Image.fromarray(np.ascontiguousarray(rgba))
    bg = Image.new("RGB", im.size, "white")
    bg.paste(im, mask=im.getchannel("A"))
    bg.thumbnail((width, width))
    buf = io.BytesIO()
    bg.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def _package(prs, title: str, thumb: bytes) -> None:
    """Document properties and preview of our own instead of python-pptx's template ones; the slide size marked
    custom (the template says 4:3)."""
    import datetime as _dt

    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    cp = prs.core_properties
    now = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0, tzinfo=None)
    cp.title, cp.author, cp.last_modified_by, cp.subject, cp.keywords, cp.comments = title, "", "", "", "", ""
    cp.created = cp.modified = now
    cp.revision = 1
    try:
        prs.part.package.part_related_by(RT.THUMBNAIL)._blob = thumb
    except KeyError:
        pass
    sld_sz = prs.part._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}sldSz")
    sld_sz.attrib.pop("type", None)


def figure_to_pptx(fig, path, dpi: float = DPI, group: bool = True) -> dict:
    """Write ``fig`` (finished, as saved) as one slide to ``path``; the figure is left as it was.  ``group``: a
    panel's picture and texts as one group.  Returns counts and remarks on what was not carried over as such."""
    from pptx import Presentation

    path = Path(path)
    w_in, h_in = fig.get_size_inches()
    axes = sorted(fig.axes, key=lambda a: a.get_zorder())  # stable: the figure's draw order
    with agg_at(fig, dpi) as canvas:
        with recording() as drawn:
            canvas.draw()
        thumb = _thumbnail(np.asarray(canvas.buffer_rgba()))
        renderer = canvas.get_renderer()
        items, remarks = texts_as_items(drawn, renderer, dpi, fig.bbox.height, axis_owners(fig), top_axes(fig))
        if drawn.as_paths:
            remarks.append(f"{len(drawn.as_paths)} texts drawn as paths (path effects) left in the pictures: "
                           + ", ".join(repr(t.get_text()) for t in drawn.as_paths[:5]))
        vis = {a: a.get_visible() for a in fig.axes}
        others = [a for a in fig.get_children() if a is not fig.patch and a not in fig.axes and a.get_visible()
                  and not isinstance(a, Text)]
        pictures: Dict[object, Tuple[bytes, Tuple[int, int, int, int]]] = {}

        def shoot(key):
            rgba = render_rgba(canvas)
            box = crop(rgba)
            if box is not None:
                x0, y0, x1, y1 = box
                pictures[key] = (png_bytes(rgba[y0:y1, x0:x1], dpi), box)
        with frozen_layout(fig), texts_hidden([t for t, _ in drawn]):
            try:
                for a in others:
                    a.set_visible(False)
                for ax in axes:
                    if vis[ax]:
                        for b in fig.axes:
                            b.set_visible(b is ax)
                        shoot(ax)
                if others:
                    for b in fig.axes:
                        b.set_visible(False)
                    for a in others:
                        a.set_visible(True)
                    shoot(None)
                    remarks.append(f"{len(others)} figure-level artists drawn as one picture above the panels")
            finally:
                for a in others:
                    a.set_visible(True)
                for a, v in vis.items():
                    a.set_visible(v)

    prs = Presentation()
    prs.slide_width, prs.slide_height = _emu(w_in * 72.0), _emu(h_in * 72.0)
    _package(prs, path.stem, thumb)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    px2pt = 72.0 / dpi
    page_w = w_in * 72.0
    counts = {"pictures": 0, "texts": 0, "boxes": 0, "groups": 0}

    def add_picture(shapes, key, name):
        png, (x0, y0, x1, y1) = pictures[key]
        pic = shapes.add_picture(io.BytesIO(png), _emu(x0 * px2pt), _emu(y0 * px2pt), _emu((x1 - x0) * px2pt),
                                 _emu((y1 - y0) * px2pt))
        pic.name = name
        pic._element.nvPicPr.cNvPr.set("descr", name)  # alt text (python-pptx writes the file name)
        counts["pictures"] += 1

    by_parent: Dict[int, List[Item]] = {}
    for it in items:
        by_parent.setdefault(id(it.parent), []).append(it)
    order = [(ax, f"panel {i + 1}") for i, ax in enumerate(axes)] + [(None, "page")]
    for key, label in order:
        members = by_parent.get(id(key), [])
        has_pic = key in pictures
        if not members and not has_pic:
            continue
        grouped = group and key is not None and (len(members) + has_pic) > 1
        shapes = slide.shapes
        if grouped:
            grp = slide.shapes.add_group_shape()
            grp.name = label
            shapes = grp.shapes
            counts["groups"] += 1
        if has_pic:
            add_picture(shapes, key, f"{label} picture")
        for it in members:
            add_item(shapes, it, _short(it.text), page_w)
            counts["texts"] += 1
            counts["boxes"] += len(it.boxes)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return {"file": str(path), **counts, "remarks": remarks}
