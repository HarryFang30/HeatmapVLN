"""PowerPoint copy of a paper figure (scripts/exp19/figures/pptx_export.py): texts as text boxes on matplotlib's
baselines, circled numbers as ovals, one picture per axes, the figure left as it was."""
from __future__ import annotations

import io

import numpy as np
import pytest

pytest.importorskip("matplotlib")
pytest.importorskip("pptx")

import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu  # noqa: E402

from scripts.exp19.figures import pptx_export as px  # noqa: E402

pytestmark = pytest.mark.filterwarnings("ignore:Glyph .* missing from font")  # the CJK test text in DejaVu

PT = 12700  # EMU


def _figure():
    fig = plt.figure(figsize=(3.0, 2.0))
    ax = fig.add_axes([0.25, 0.15, 0.45, 0.45])
    ax.imshow(np.random.RandomState(0).rand(8, 8), cmap="Greys")
    ax.set_xticks([0, 4])
    ax.set_yticks([])
    ax.text(2, 2, "1", ha="center", va="center", fontsize=6, fontweight="bold", color="white",
            bbox=dict(boxstyle="circle,pad=0.14", fc="black", ec="white", lw=0.5))
    fig.text(0.05, 0.92, "Left", ha="left", fontsize=7)
    fig.text(0.95, 0.92, "Right italic", ha="right", fontsize=6, fontstyle="italic")
    fig.text(0.06, 0.35, "Rotated", rotation=90, ha="center", va="center", fontsize=6)
    fig.text(0.85, 0.35, "two\nlines", ha="center", va="center", fontsize=6, linespacing=2.0)
    px.same_box(fig.text(0.05, 0.75, "(a) Title", fontsize=7, fontweight="bold"), "head")
    px.same_box(fig.text(0.40, 0.75, "Success", fontsize=6.3, color="0.4"), "head")
    fig.text(0.75, 0.05, "中文 ab", fontsize=6)
    return fig


def _rgba(fig, dpi=200):
    buf = io.BytesIO()
    fig.savefig(buf, format="rgba", dpi=dpi)
    return buf.getvalue()


def _shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from _shapes(sh.shapes)
        else:
            yield sh


def _recorded(fig):
    """(text, first-line baseline-start x, y in pt from the top-left) of every drawn line, at the export dpi."""
    with px.agg_at(fig, px.DPI) as canvas:
        with px.recording() as drawn:
            canvas.draw()
    k = 72.0 / px.DPI
    return {t.get_text(): [(s, x * k, y * k, w * k) for x, y, s, _, _, _, w in lines] for t, lines in drawn}


@pytest.fixture(scope="module")
def exported(tmp_path_factory):
    fig = _figure()
    before = _rgba(fig)
    rec = _recorded(fig)
    path = tmp_path_factory.mktemp("pptx") / "fig.pptx"
    state = (fig.dpi, fig.canvas, fig.patch.get_visible())
    info = px.figure_to_pptx(fig, path)
    restored = (fig.dpi, fig.canvas, fig.patch.get_visible()) == state
    after = _rgba(fig)
    prs = Presentation(str(path))
    shapes = list(_shapes(prs.slides[0].shapes))
    texts = {sh.text_frame.text: sh for sh in shapes if sh.has_text_frame and sh.text_frame.text}
    yield {"fig": fig, "prs": prs, "info": info, "shapes": shapes, "texts": texts, "rec": rec,
           "same": before == after and restored, "top": prs.slides[0].shapes}
    plt.close(fig)


def _top(sh) -> float:
    return Emu(sh.top).pt


def test_slide_is_the_figure_and_the_figure_is_left_as_it_was(exported):
    prs, fig = exported["prs"], exported["fig"]
    w, h = fig.get_size_inches()
    assert prs.slide_width == round(w * 72 * PT) and prs.slide_height == round(h * 72 * PT)
    assert exported["same"]  # texts, axes, background, dpi and canvas restored
    assert exported["info"]["remarks"] == []
    prs = exported["prs"]
    cp = prs.core_properties
    assert cp.title == "fig" and cp.author == "" and cp.last_modified_by == ""  # not the template's
    assert prs.part._element.find(qn("p:sldSz")).get("type") is None  # custom size, not "screen4x3"
    from PIL import Image
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    thumb = Image.open(io.BytesIO(prs.part.package.part_related_by(RT.THUMBNAIL).blob))
    assert thumb.size[0] / thumb.size[1] == pytest.approx(1.5, abs=0.02)  # a preview of this 3 x 2 in figure


def test_every_text_is_a_text_box_and_the_heading_runs_one(exported):
    texts = exported["texts"]
    for want in ("Left", "Right italic", "Rotated", "two\nlines", "中文 ab", "1", "0", "4"):
        assert want in texts, (want, sorted(texts))
    head = texts["(a) Title Success"]  # one box (``same_box``): title, a widened space, outcome
    runs = head.text_frame.paragraphs[0].runs
    assert [r.text for r in runs] == ["(a) Title", " ", "Success"] and runs[0].font.bold and not runs[2].font.bold
    rec = exported["rec"]
    gap = rec["Success"][0][1] - (rec["(a) Title"][0][1] + rec["(a) Title"][0][3])
    spc = int(runs[1]._r.rPr.get("spc")) / 100.0
    assert spc == pytest.approx(gap - px.SPACE_EM * 6.3, abs=0.011)


@pytest.mark.parametrize("name", ["Left", "Right italic", "(a) Title Success", "中文 ab", "1"])
def test_one_line_sits_on_matplotlibs_baseline_and_start(exported, name):
    sh, rec = exported["texts"][name], exported["rec"]
    first = name.split(" Success")[0]
    _, x, y, w = rec[first][0]
    p = sh.text_frame.paragraphs[0]
    size = max(r.font.size.pt for r in p.runs)
    cjk = any(px._CJK.search(r.text) for r in p.runs)
    asc = px.METRICS[px.CJK_FONT if cjk else px.LATIN_FONT][0] * size
    lead = int(p._p.pPr.find(qn("a:lnSpc"))[0].get("val")) / 100.0
    desc = px.METRICS[px.CJK_FONT if cjk else px.LATIN_FONT][1] * size
    assert lead == pytest.approx(asc + desc, abs=0.011)  # "exactly" ascent + descent
    assert _top(sh) + asc == pytest.approx(y, abs=0.02)  # baseline one ascent below the top
    algn = p._p.pPr.get("algn")
    left, width = Emu(sh.left).pt, Emu(sh.width).pt
    if algn == "l":
        assert left == pytest.approx(x, abs=0.02)
    elif algn == "r":
        assert left + width == pytest.approx(x + w, abs=0.02)
    else:
        assert left + width / 2 == pytest.approx(x + w / 2, abs=0.02)
    assert width > w  # room to spare, so an editor that wraps anyway keeps the line whole
    assert sh.text_frame._bodyPr.get("wrap") == "none"


def test_rotated_text_turns_about_its_centre(exported):
    sh = exported["texts"]["Rotated"]
    assert sh.rotation == pytest.approx(270.0)
    _, x, y, w = exported["rec"]["Rotated"][0]  # baseline start: the line runs upwards from (x, y)
    asc = px.METRICS[px.LATIN_FONT][0] * 6
    desc = px.METRICS[px.LATIN_FONT][1] * 6
    cx, cy = Emu(sh.left).pt + Emu(sh.width).pt / 2, Emu(sh.top).pt + Emu(sh.height).pt / 2
    assert cy == pytest.approx(y - w / 2, abs=0.02)  # centred along the line
    assert cx == pytest.approx(x - (asc - desc) / 2, abs=0.02)  # the line box across the baseline


def test_two_lines_keep_their_pitch(exported):
    sh, rec = exported["texts"]["two\nlines"], exported["rec"]["two\nlines"]
    ps = sh.text_frame.paragraphs
    assert len(ps) == 2
    asc, desc = px.METRICS[px.LATIN_FONT][0] * 6, px.METRICS[px.LATIN_FONT][1] * 6
    before = int(ps[1]._p.pPr.find(qn("a:spcBef"))[0].get("val")) / 100.0
    assert _top(sh) + asc == pytest.approx(rec[0][2], abs=0.02)
    assert _top(sh) + asc + desc + before + asc == pytest.approx(rec[1][2], abs=0.02)


def test_cjk_text_gets_yahei_and_chinese_language(exported):
    r = exported["texts"]["中文 ab"].text_frame.paragraphs[0].runs[0]._r.rPr
    assert r.get("lang") == "zh-CN" and r.find(qn("a:ea")).get("typeface") == px.CJK_FONT
    assert r.find(qn("a:latin")).get("typeface") == px.LATIN_FONT


def test_circled_number_is_an_oval_right_under_its_text(exported):
    shapes = exported["shapes"]
    i = next(i for i, sh in enumerate(shapes) if sh.has_text_frame and sh.text_frame.text == "1")
    oval = shapes[i - 1]
    geom = oval._element.spPr.find(qn("a:prstGeom")).get("prst")
    assert geom == "ellipse" and oval.fill.fore_color.rgb == RGBColor.from_string("000000")
    assert oval.line.color.rgb == RGBColor.from_string("FFFFFF") and oval.line.width.pt == pytest.approx(0.5)
    assert oval._element.find(qn("p:style")) is None  # no theme line or effects
    num = shapes[i]
    c_o = (Emu(oval.left).pt + Emu(oval.width).pt / 2, Emu(oval.top).pt + Emu(oval.height).pt / 2)
    _, x, y, w = exported["rec"]["1"][0]
    assert c_o[0] == pytest.approx(x + w / 2, abs=0.6)  # the number is centred in its circle
    assert Emu(num.left).pt + Emu(num.width).pt / 2 == pytest.approx(x + w / 2, abs=0.02)


def test_one_picture_per_axes_with_marks(exported):
    fig, prs = exported["fig"], exported["prs"]
    pics = [sh for sh in exported["shapes"] if sh.shape_type == 13]
    assert len(pics) == exported["info"]["pictures"] == 1  # the image axes; no figure-level marks
    ax = fig.axes[0].get_position()
    w, h = fig.get_size_inches() * 72
    pic = pics[0]
    left, top = Emu(pic.left).pt, Emu(pic.top).pt
    assert left <= ax.x0 * w + 0.01 and top <= (1 - ax.y1) * h + 0.01  # the axes' frame and tick marks inside
    assert left + Emu(pic.width).pt >= ax.x1 * w - 0.01
    groups = [sh for sh in prs.slides[0].shapes if sh.shape_type == 6]
    assert len(groups) == 1  # the tick labels (no ``.axes`` of their own) move with their panel
    assert {s.name for s in groups[0].shapes} == {"panel 1 picture", "1", "box 1", "0", "4"}


def test_texts_merge_only_on_a_shared_key_and_svg_ids_stay_unique(tmp_path):
    fig = plt.figure(figsize=(2.0, 1.0))
    fig.text(0.1, 0.5, "a", fontsize=6)
    fig.text(0.3, 0.5, "b", fontsize=6)
    px.same_box(fig.text(0.1, 0.2, "c", fontsize=6), "k")
    px.same_box(fig.text(0.3, 0.2, "d", fontsize=6), "k")
    svg = io.BytesIO()
    fig.savefig(svg, format="svg")
    assert b"_pptx_box" not in svg.getvalue() and b'id="k"' not in svg.getvalue()
    px.figure_to_pptx(fig, tmp_path / "x.pptx")
    plt.close(fig)
    texts = sorted(sh.text_frame.text for sh in Presentation(str(tmp_path / "x.pptx")).slides[0].shapes)
    assert texts == ["a", "b", "c d"]


def _baselines_centred(sh):
    """Baselines (pt from the page top) of a text box's lines where line boxes are centred on the glyphs."""
    out, top = [], Emu(sh.top).pt
    for p in sh.text_frame.paragraphs:
        L = int(p._p.pPr.find(qn("a:lnSpc"))[0].get("val")) / 100.0
        top += int(p._p.pPr.find(qn("a:spcBef"))[0].get("val")) / 100.0
        runs = [px.Run(r.text, r.font.size.pt, False, False, (0, 0, 0, 1), 0, 0, 0) for r in p.runs if r.text.strip()]
        a, d = px.line_metrics(runs)
        out.append(top + L / 2 + (a - d) / 2)
        top += L
    return out


def test_close_lines_keep_their_baselines(tmp_path):
    """A CJK title over a Latin line closer than descent + ascent (figure B's zh headings), and stacked CJK
    characters: every baseline where matplotlib drew it, with the line boxes centred on the glyphs."""
    fig = plt.figure(figsize=(3.0, 1.5))
    px.same_box(fig.text(0.05, 0.6, "标题 Title", fontsize=7.5), "h")
    px.same_box(fig.text(0.05, 0.6 - 6.8 / 108.0, "second line", fontsize=6), "h")
    fig.text(0.8, 0.5, "历\n史", fontsize=6, linespacing=1.05, ha="center", va="center")
    rec = _recorded(fig)
    px.figure_to_pptx(fig, tmp_path / "c.pptx")
    plt.close(fig)
    texts = {sh.text_frame.text: sh for sh in Presentation(str(tmp_path / "c.pptx")).slides[0].shapes
             if sh.has_text_frame}
    head = texts["标题 Title\nsecond line"]
    want = [rec["标题 Title"][0][2], rec["second line"][0][2]]
    assert _baselines_centred(head) == pytest.approx(want, abs=0.02)
    stacked = texts["历\n史"]
    assert _baselines_centred(stacked) == pytest.approx([ln[2] for ln in rec["历\n史"]], abs=0.02)
    for sh in (head, stacked):  # tight: line heights below ascent + descent
        assert int(sh.text_frame.paragraphs[1]._p.pPr.find(qn("a:lnSpc"))[0].get("val")) / 100.0 < 7.0


def test_annotation_arrows_inset_axes_and_box_styles(tmp_path):
    fig = plt.figure(figsize=(3.0, 2.0))
    ax = fig.add_axes([0.1, 0.1, 0.8, 0.8])
    ax.set_axis_off()
    ins = ax.inset_axes([0.6, 0.6, 0.3, 0.3])
    ins.set_xticks([])
    ins.set_yticks([])
    ins.text(0.5, 0.5, "inset", ha="center", fontsize=6)
    ax.annotate("note", (0.2, 0.2), xytext=(0.5, 0.3), arrowprops=dict(arrowstyle="-", lw=2), fontsize=6)
    for i, style in enumerate(("ellipse,pad=0.2", "round,pad=0.2,rounding_size=0", "round,pad=0.2")):
        ax.text(0.1 + 0.15 * i, 0.8, str(i), fontsize=6, bbox=dict(boxstyle=style, fc="k"), color="w")
    info = px.figure_to_pptx(fig, tmp_path / "a.pptx")
    plt.close(fig)
    prs = Presentation(str(tmp_path / "a.pptx"))
    (grp,) = [sh for sh in prs.slides[0].shapes if sh.shape_type == 6]
    names = [sh.name for sh in grp.shapes]
    assert "inset" in names and "note" in names  # the inset's text goes with the panel that draws it
    geoms = {sh.name: sh._element.spPr.find(qn("a:prstGeom")).get("prst") for sh in grp.shapes
             if sh.name.startswith("box ")}
    assert geoms == {"box 0": "ellipse", "box 1": "rect", "box 2": "roundRect"}
    pic = next(sh for sh in grp.shapes if sh.shape_type == 13)
    from PIL import Image

    alpha = np.asarray(Image.open(io.BytesIO(pic.image.blob)))[..., 3]
    x0, y0 = Emu(pic.left).pt, Emu(pic.top).pt
    k = alpha.shape[1] / Emu(pic.width).pt  # px per pt
    mid = (0.1 + 0.8 * 0.35) * 216 - x0, (1 - (0.1 + 0.8 * 0.25)) * 144 - y0  # the arrow's midpoint (pt)
    assert alpha[int(mid[1] * k) - 3: int(mid[1] * k) + 4, int(mid[0] * k) - 3: int(mid[0] * k) + 4].max() > 0
    assert info["remarks"] == []


def test_layout_engine_positions_hold_and_the_engine_comes_back(tmp_path):
    fig = plt.figure(figsize=(3.0, 2.0), layout="constrained")
    axs = fig.subplots(1, 2)
    for a in axs:
        a.set_title("A long title", fontsize=6)
        a.set_ylabel("y label", fontsize=6)
    engine = fig.get_layout_engine()
    px.figure_to_pptx(fig, tmp_path / "l.pptx")
    assert fig.get_layout_engine() is engine
    pos = [a.get_position() for a in axs]  # as the export's first draw left them
    plt.close(fig)
    prs = Presentation(str(tmp_path / "l.pptx"))
    pics = sorted((sh for g in prs.slides[0].shapes if g.shape_type == 6 for sh in g.shapes if sh.shape_type == 13),
                  key=lambda sh: sh.left)
    assert len(pics) == 2
    for pic, bb in zip(pics, pos):  # each picture: its axes frame (and tick marks) where the texts were placed
        left, right = Emu(pic.left).pt, Emu(pic.left + pic.width).pt
        assert bb.x0 * 216 - 5.0 <= left <= bb.x0 * 216 + 0.5
        assert bb.x1 * 216 - 0.5 <= right <= bb.x1 * 216 + 1.0


def test_box_slack_stays_on_the_page(tmp_path):
    fig = plt.figure(figsize=(2.0, 1.0))
    fig.text(0.02, 0.5, "a line that runs to the right edge", fontsize=6)
    px.figure_to_pptx(fig, tmp_path / "w.pptx")
    plt.close(fig)
    prs = Presentation(str(tmp_path / "w.pptx"))
    (sh,) = list(prs.slides[0].shapes)
    assert sh.left + sh.width <= prs.slide_width
